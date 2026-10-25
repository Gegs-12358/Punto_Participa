"""
Lógica del importador de alumnos (.xlsx / .csv).

Separada del comando para reutilizarla en la vista web:
    filas = leer_filas(archivo, nombre)      # lectura + validación de contenido
    resultado = analizar(filas)              # simulación: NO escribe en la BBDD
    aplicar(resultado)                       # escribe con save() (limpia campos)

Reglas:
- La clave es el documento normalizado (normalizar_documento).
- Existe -> se compara campo por campo y se actualiza lo que cambió.
- No existe -> se crea.
- Está en la BBDD (activo) y NO viene en el archivo -> se BLOQUEA (activo=False).
  Nunca se borra: inscripciones y asistencias quedan intactas.
- Si un bloqueado vuelve a venir en el archivo, se reactiva.
- Una fila con error NO bloquea al alumno: su documento sí viene en el archivo.
- Los errores son por fila y no tumban el archivo.
- Altas y cambios usan save() (respeta Alumno.save()). El bloqueo usa update()
  solo sobre `activo`, que no necesita limpieza.
"""

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction

from .models import Alumno, Carrera, Jornada
from .utils import (
    limpiar_texto,
    normalizar_documento,
    validar_documento,
    validar_rut,
    validar_rut_extranjero,
)

MAX_FILAS = 5000
MAX_BYTES = 5 * 1024 * 1024

COLUMNAS_OBLIGATORIAS = ['rut', 'nombres', 'apellidos', 'correo', 'carrera', 'jornada']

ALIAS_COLUMNAS = {
    'run': 'rut', 'documento': 'rut', 'rut_pasaporte': 'rut',
    'numero_documento': 'rut', 'identificacion': 'rut',
    'tipo': 'tipo_documento', 'tipo_de_documento': 'tipo_documento',
    'nombre': 'nombres',
    'apellido': 'apellidos',
    'email': 'correo', 'correo_electronico': 'correo', 'mail': 'correo',
}

TIPOS_VALIDOS = {t for t, _ in Alumno.TIPO_DOCUMENTO_CHOICES}


class ArchivoInvalido(Exception):
    """El archivo completo no se puede procesar (formato, tamaño, columnas)."""


@dataclass
class Resultado:
    nuevos: list = field(default_factory=list)        # [(fila, datos)]
    actualizados: list = field(default_factory=list)  # [(fila, datos, [(campo, antes, despues)])]
    bloqueados: list = field(default_factory=list)    # [(documento, nombre)]
    total_activos: int = 0                            # activos en la BBDD antes de importar
    sin_cambios: int = 0
    errores: list = field(default_factory=list)       # [(fila, mensaje)]
    avisos: list = field(default_factory=list)        # [(fila, mensaje)]

    @property
    def total(self):
        return len(self.nuevos) + len(self.actualizados) + self.sin_cambios + len(self.errores)


# ------------------------------------------------------------
# Lectura
# ------------------------------------------------------------

def _clave_columna(encabezado):
    h = limpiar_texto(encabezado) or ''
    h = unicodedata.normalize('NFKD', h).encode('ascii', 'ignore').decode().lower()
    h = re.sub(r'[^a-z0-9]+', '_', h).strip('_')
    return ALIAS_COLUMNAS.get(h, h)


def _celda_a_texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor)


def _filas_csv(contenido):
    try:
        texto = contenido.decode('utf-8-sig')
    except UnicodeDecodeError:
        texto = contenido.decode('latin-1')
    primera = texto.split('\n', 1)[0]
    delimitador = max([';', ',', '\t'], key=primera.count)
    return list(csv.reader(io.StringIO(texto), delimiter=delimitador))


def _filas_xlsx(contenido):
    import zipfile
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError):
        raise ArchivoInvalido('El archivo no es un .xlsx válido.')
    ws = wb.active
    return [[_celda_a_texto(c) for c in fila] for fila in ws.iter_rows(values_only=True)]


def leer_filas(archivo, nombre):
    """Devuelve [(numero_fila, {columna: texto})]. `archivo` es un objeto con .read()."""
    contenido = archivo.read()
    if len(contenido) > MAX_BYTES:
        raise ArchivoInvalido(f'El archivo supera {MAX_BYTES // (1024 * 1024)} MB.')

    nombre = (nombre or '').lower()
    if nombre.endswith('.xlsx'):
        if not contenido.startswith(b'PK'):          # contenido, no solo extensión
            raise ArchivoInvalido('El archivo no es un .xlsx válido.')
        matriz = _filas_xlsx(contenido)
    elif nombre.endswith('.csv'):
        if b'\x00' in contenido[:4096]:
            raise ArchivoInvalido('El archivo no es un CSV de texto válido.')
        matriz = _filas_csv(contenido)
    else:
        raise ArchivoInvalido('Formato no permitido. Usa .xlsx o .csv.')

    matriz = [f for f in matriz if any(_celda_a_texto(c).strip() for c in f)]
    if not matriz:
        raise ArchivoInvalido('El archivo está vacío.')

    columnas = [_clave_columna(c) for c in matriz[0]]
    faltan = [c for c in COLUMNAS_OBLIGATORIAS if c not in columnas]
    if faltan:
        raise ArchivoInvalido('Faltan columnas: ' + ', '.join(faltan))

    cuerpo = matriz[1:]
    if len(cuerpo) > MAX_FILAS:
        raise ArchivoInvalido(f'Demasiadas filas (máximo {MAX_FILAS}).')

    filas = []
    for i, fila in enumerate(cuerpo, start=2):       # la fila 1 es el encabezado
        datos = {}
        for col, valor in zip(columnas, fila):
            datos.setdefault(col, _celda_a_texto(valor))
        filas.append((i, datos))
    return filas


# ------------------------------------------------------------
# Análisis (no escribe)
# ------------------------------------------------------------

def inferir_tipo(doc):
    if validar_rut(doc):
        return 'RUT'
    if validar_rut_extranjero(doc):
        return 'RUN_PROVISORIO'
    return 'PASAPORTE'


def _normalizar_tipo(texto):
    t = limpiar_texto(texto).upper().replace(' ', '_')
    return {'CEDULA': 'CEDULA_EXTRANJERO', 'RUN': 'RUT'}.get(t, t)


def analizar(filas):
    res = Resultado()
    carreras = {limpiar_texto(c.nombre).casefold(): c.nombre for c in Carrera.objects.all()}
    jornadas = {limpiar_texto(j.nombre).casefold(): j.nombre for j in Jornada.objects.all()}

    docs = {normalizar_documento(d.get('rut')) for _, d in filas}
    existentes = {a.rut: a for a in Alumno.objects.filter(rut__in=[d for d in docs if d])}
    res.total_activos = Alumno.objects.filter(activo=True).count()

    max_len = {f: Alumno._meta.get_field(f).max_length
               for f in ('nombres', 'apellidos', 'carrera', 'jornada')}
    vistos = {}

    for n, d in filas:
        errores = []
        doc = normalizar_documento(d.get('rut'))
        if not doc:
            res.errores.append((n, 'Falta el documento.'))
            continue
        if doc in vistos:
            res.errores.append((n, f'Documento repetido en el archivo (ver fila {vistos[doc]}).'))
            continue
        vistos[doc] = n

        aviso = None
        tipo_txt = d.get('tipo_documento', '')
        if limpiar_texto(tipo_txt):
            tipo = _normalizar_tipo(tipo_txt)
            if tipo not in TIPOS_VALIDOS:
                errores.append(f'Tipo de documento desconocido: "{limpiar_texto(tipo_txt)}".')
        else:
            tipo = inferir_tipo(doc)
            aviso = f'Tipo de documento inferido: {tipo}.'
        if not errores and not validar_documento(doc, tipo):
            errores.append(f'Documento no válido para el tipo {tipo}.')

        nombres = limpiar_texto(d.get('nombres'))
        apellidos = limpiar_texto(d.get('apellidos'))
        correo = limpiar_texto(d.get('correo'))
        if not nombres:
            errores.append('Falta el nombre.')
        if not apellidos:
            errores.append('Falta el apellido.')
        try:
            validate_email(correo)
        except ValidationError:
            errores.append('Correo no válido.')

        carrera = carreras.get(limpiar_texto(d.get('carrera')).casefold())
        if not carrera:
            errores.append(f'Carrera desconocida: "{limpiar_texto(d.get("carrera"))}".')
        jornada = jornadas.get(limpiar_texto(d.get('jornada')).casefold())
        if not jornada:
            errores.append(f'Jornada desconocida: "{limpiar_texto(d.get("jornada"))}".')

        for campo, valor in (('nombres', nombres), ('apellidos', apellidos)):
            if valor and len(valor) > max_len[campo]:
                errores.append(f'{campo.capitalize()} supera {max_len[campo]} caracteres.')

        if errores:
            res.errores.append((n, ' '.join(errores)))
            continue
        if aviso:
            res.avisos.append((n, aviso))

        datos = dict(rut=doc, tipo_documento=tipo, nombres=nombres, apellidos=apellidos,
                     correo=correo, carrera=carrera, jornada=jornada, activo=True)
        actual = existentes.get(doc)
        if actual is None:
            res.nuevos.append((n, datos))
        else:
            cambios = [(c, getattr(actual, c), v) for c, v in datos.items()
                       if getattr(actual, c) != v]
            if cambios:
                res.actualizados.append((n, datos, cambios))
            else:
                res.sin_cambios += 1

    # Alumnos activos que no vienen en el archivo -> se bloquean.
    # `docs` incluye también las filas con error: esas NO se bloquean.
    ausentes = (Alumno.objects.filter(activo=True, rut__isnull=False)
                .exclude(rut__in=[d for d in docs if d]))
    res.bloqueados = [(a.rut, f'{a.nombres} {a.apellidos}') for a in ausentes]
    return res


def bloqueo_masivo(res, umbral=0.2):
    """True si se bloquearía más del 20 % de los activos: casi siempre es un
    archivo equivocado o incompleto, no una baja real."""
    return res.total_activos > 0 and len(res.bloqueados) > umbral * res.total_activos


# ------------------------------------------------------------
# Aplicación (escribe)
# ------------------------------------------------------------

@transaction.atomic
def aplicar(res):
    """Crea, actualiza y bloquea. Todo o nada. Devuelve (creados, actualizados, bloqueados)."""
    for _, datos in res.nuevos:
        Alumno(**datos).save()
    for _, datos, _cambios in res.actualizados:
        alumno = Alumno.objects.get(rut=datos['rut'])
        for campo, valor in datos.items():
            setattr(alumno, campo, valor)
        alumno.save()
    docs_bloqueo = [doc for doc, _ in res.bloqueados]
    if docs_bloqueo:
        Alumno.objects.filter(rut__in=docs_bloqueo).update(activo=False)
    return len(res.nuevos), len(res.actualizados), len(docs_bloqueo)
