"""
Pruebas de exportación CSV/Excel con nombres con tildes, ñ, apóstrofes,
comas, saltos de línea, fórmulas y caracteres de control.

Se guarda como gestion/tests_exportaciones.py (Django lo descubre solo).
No usa base de datos: prueba directamente las funciones de exportación.
"""

import csv
import io

from django.test import SimpleTestCase
from openpyxl import load_workbook

from .views import _exportar_csv, _exportar_excel, _sanitizar_celda

NOMBRES = [
    "Peña Muñoz",
    "Ñancupil Müller-Åström",
    "Nguyễn Văn Ánh",
    "Bernardo O`higgins",
    "Bernardo O'Higgins",
    'María "La Pepa" Pérez',
    "100% & más #1",
    "😀 Emoji",
    "Pérez, Juan",
    "Salto\nde línea",
]
FORMULAS = [
    '=HYPERLINK("http://x.com")',
    "+56912345678",
    "-Pedro",
    "@usuario",
]
TODOS = NOMBRES + FORMULAS


class SanitizarCeldaTestCase(SimpleTestCase):

    def test_no_altera_nombres_legitimos(self):
        for n in NOMBRES:
            with self.subTest(n=n):
                self.assertEqual(_sanitizar_celda(n), n)

    def test_numeros_pasan_sin_cambios(self):
        self.assertEqual(_sanitizar_celda(-5), -5)
        self.assertEqual(_sanitizar_celda(3.5), 3.5)

    def test_neutraliza_formulas(self):
        for f in FORMULAS:
            with self.subTest(f=f):
                self.assertEqual(_sanitizar_celda(f), "'" + f)

    def test_quita_caracteres_de_control_y_luego_revisa_formula(self):
        self.assertEqual(_sanitizar_celda("A\x0bB"), "AB")
        self.assertEqual(_sanitizar_celda("x\x00y"), "xy")
        # el NUL escondido antes de '=' no debe burlar la protección
        self.assertEqual(_sanitizar_celda("\x00=cmd"), "'=cmd")


class ExportarCsvTestCase(SimpleTestCase):

    def test_csv_incluye_bom_para_excel(self):
        r = _exportar_csv([[n] for n in TODOS], ['Nombre'], 'x')
        self.assertTrue(r.content.startswith(b'\xef\xbb\xbf'))
        self.assertIn('charset=utf-8', r['Content-Type'])

    def test_csv_conserva_nombres(self):
        r = _exportar_csv([[n] for n in TODOS], ['Nombre'], 'x')
        texto = r.content.decode('utf-8-sig')
        filas = list(csv.reader(io.StringIO(texto)))
        self.assertEqual(filas[0], ['Nombre'])
        valores = [f[0] for f in filas[1:]]
        self.assertEqual(valores, [_sanitizar_celda(n) for n in TODOS])


class ExportarExcelTestCase(SimpleTestCase):

    def _leer(self, respuesta, n_filas):
        ws = load_workbook(io.BytesIO(respuesta.content)).active
        return ws, [ws.cell(row=i, column=1).value for i in range(2, 2 + n_filas)]

    def test_excel_conserva_nombres_y_no_crea_formulas(self):
        r = _exportar_excel([[n] for n in TODOS], ['Nombre'], 'x')
        ws, valores = self._leer(r, len(TODOS))
        self.assertEqual(valores, [_sanitizar_celda(n) for n in TODOS])
        for i in range(2, 2 + len(TODOS)):
            self.assertNotEqual(ws.cell(row=i, column=1).data_type, 'f')

    def test_excel_con_caracteres_ilegales_no_revienta(self):
        r = _exportar_excel([["A\x0bB", "x\x00y"]], ['a', 'b'], 'x')
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws['A2'].value, 'AB')
        self.assertEqual(ws['B2'].value, 'xy')
