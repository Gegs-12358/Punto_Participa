"""
Suite complementaria de casos extremos y seguridad para Punto Participa.

No requiere Docker. Se ejecuta contra la base de datos temporal de Django:
    python manage.py test gestion.tests_extremos -v 2

Incluye validaciones de contraseñas, autenticación, restricciones de datos,
creación AJAX, correos fallidos, límites de invitaciones y casos inválidos.
"""

from datetime import timedelta
from io import BytesIO

from django.contrib.auth import authenticate
from django.core import mail
from django.core.exceptions import ValidationError
from django.core.mail.backends.base import BaseEmailBackend
from django.db import IntegrityError, transaction
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from openpyxl import load_workbook

from .forms import ActividadForm
from .models import Actividad, Asistencia, Inscripcion, LogAuditoria, NotificacionCorreo
from .tests_integrales import BaseSistemaTestCase


class BackendCorreoFallido(BaseEmailBackend):
    """Backend controlado para probar el registro de correos fallidos."""

    def send_messages(self, email_messages):
        raise OSError("Servidor SMTP de prueba no disponible")


class ContraseñasYAutenticacionExtremaTests(BaseSistemaTestCase):
    def test_cambio_contrasena_incorrecta_no_modifica_la_cuenta(self):
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("cambiar_contrasena"),
            {
                "contrasena_actual": "incorrecta",
                "nueva_contrasena": "NuevaClave123!",
                "confirmar_contrasena": "NuevaClave123!",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(authenticate(username=self.admin.username, password="ClaveSegura123!"))
        self.assertFalse(authenticate(username=self.admin.username, password="NuevaClave123!"))

    def test_cambio_contrasena_con_confirmacion_distinta_no_modifica(self):
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("cambiar_contrasena"),
            {
                "contrasena_actual": "ClaveSegura123!",
                "nueva_contrasena": "NuevaClave123!",
                "confirmar_contrasena": "OtraClave123!",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(authenticate(username=self.admin.username, password="ClaveSegura123!"))

    def test_cambio_contrasena_valido_actualiza_password_y_perfil(self):
        self.admin.usuariosistema.must_change_password = True
        self.admin.usuariosistema.save(update_fields=["must_change_password"])
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("cambiar_contrasena"),
            {
                "contrasena_actual": "ClaveSegura123!",
                "nueva_contrasena": "NuevaClave123!",
                "confirmar_contrasena": "NuevaClave123!",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("dashboard"))
        self.admin.refresh_from_db()
        self.admin.usuariosistema.refresh_from_db()
        self.assertTrue(self.admin.check_password("NuevaClave123!"))
        self.assertFalse(self.admin.usuariosistema.must_change_password)
        self.assertTrue(LogAuditoria.objects.filter(accion="Cambio de contraseña").exists())

    def test_login_con_usuario_inactivo_no_autentica(self):
        self.admin.usuariosistema.activo = False
        self.admin.usuariosistema.save(update_fields=["activo"])
        response = self.client.post(
            reverse("login"),
            {"username": self.admin.username, "password": "ClaveSegura123!"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.wsgi_request.user.is_authenticated)

    def test_login_respeta_next_interno(self):
        response = self.client.post(
            reverse("login") + "?next=/escaneo/",
            {
                "username": self.admin.username,
                "password": "ClaveSegura123!",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, "/escaneo/")

    def test_login_ignora_next_externo(self):
        response = self.client.post(
            reverse("login") + "?next=https://sitio-malicioso.example/",
            {
                "username": self.admin.username,
                "password": "ClaveSegura123!",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("dashboard"))

    def test_logout_cierra_la_sesion(self):
        self.client.force_login(self.admin)
        response = self.client.post(reverse("logout"))
        self.assertEqual(response.status_code, 302)
        protegida = self.client.get(reverse("reportes"))
        self.assertEqual(protegida.status_code, 302)


class RestriccionesDeDatosExtremasTests(BaseSistemaTestCase):
    def test_inscripcion_duplicada_es_rechazada_por_la_base(self):
        Inscripcion.objects.create(actividad=self.actividad, alumno=self.alumno)
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Inscripcion.objects.create(actividad=self.actividad, alumno=self.alumno)

    def test_asistencia_duplicada_es_rechazada_por_la_base(self):
        Asistencia.objects.create(
            actividad=self.actividad_masiva,
            alumno=self.alumno,
            metodo_ingreso="RUT",
            registrado_por=self.registrador,
        )
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Asistencia.objects.create(
                    actividad=self.actividad_masiva,
                    alumno=self.alumno,
                    metodo_ingreso="QR",
                    registrado_por=self.registrador,
                )

    def test_eliminar_actividad_cascadea_inscripciones_y_asistencias(self):
        inscripcion = Inscripcion.objects.create(
            actividad=self.actividad, alumno=self.alumno
        )
        asistencia = Asistencia.objects.create(
            actividad=self.actividad,
            alumno=self.alumno,
            metodo_ingreso="RUT",
            registrado_por=self.registrador,
        )
        actividad_id = self.actividad.pk
        self.actividad.delete()
        self.assertFalse(Actividad.objects.filter(pk=actividad_id).exists())
        self.assertFalse(Inscripcion.objects.filter(pk=inscripcion.pk).exists())
        self.assertFalse(Asistencia.objects.filter(pk=asistencia.pk).exists())

    def test_recalculo_de_cupos_nunca_queda_negativo(self):
        self.actividad.cupos_totales = 1
        self.actividad.recalcular_cupos(inscritos_override=99)
        self.assertEqual(self.actividad.cupos_disponibles, 0)
        self.assertTrue(self.actividad.esta_lleno)

    def test_formulario_limpia_cupos_de_actividad_masiva(self):
        inicio = timezone.localtime(timezone.now() + timedelta(days=4))
        form = ActividadForm(
            data={
                "titulo": "Masiva sin cupos",
                "descripcion": "Actividad general",
                "tipo": "MASIVA",
                "lugar": "Auditorio",
                "cupos_totales": 500,
                "fecha_inicio": inicio.strftime("%Y-%m-%dT%H:%M"),
                "fecha_fin": (inicio + timedelta(hours=2)).strftime("%Y-%m-%dT%H:%M"),
                "carreras": [self.carrera.pk],
                "jornadas": [self.jornada.pk],
            }
        )
        self.assertTrue(form.is_valid(), form.errors.as_json())
        actividad = form.save(commit=False)
        self.assertIsNone(actividad.cupos_totales)
        self.assertIsNone(actividad.cupos_disponibles)


class VistasAjaxYValidacionesExtremasTests(BaseSistemaTestCase):
    def datos_creacion(self, titulo="Actividad AJAX válida"):
        inicio = timezone.localtime(timezone.now() + timedelta(days=6))
        return {
            "titulo": titulo,
            "descripcion": "Descripción de una actividad nueva",
            "tipo": "TALLER",
            "lugar": "Sala AJAX",
            "cupos_totales": 10,
            "fecha_inicio": inicio.strftime("%Y-%m-%dT%H:%M"),
            "fecha_fin": (inicio + timedelta(hours=2)).strftime("%Y-%m-%dT%H:%M"),
            "carreras": [self.carrera.pk],
            "jornadas": [self.jornada.pk],
        }

    def test_crear_actividad_ajax_valida_devuelve_id(self):
        self.client.force_login(self.creador)
        response = self.client.post(
            reverse("crear_actividad"),
            self.datos_creacion(),
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200)
        datos = response.json()
        self.assertTrue(datos["success"])
        self.assertTrue(Actividad.objects.filter(pk=datos["id"]).exists())

    def test_crear_actividad_ajax_invalida_devuelve_errores(self):
        self.client.force_login(self.creador)
        response = self.client.post(
            reverse("crear_actividad"),
            {"titulo": ""},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json()["success"])
        self.assertIn("formActividad", response.json()["html"])

    def test_crear_actividad_get_partial_devuelve_fragmento(self):
        self.client.force_login(self.creador)
        response = self.client.get(reverse("crear_actividad") + "?partial=1")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="formActividad"')

    def test_crear_actividad_sin_rol_es_rechazado(self):
        self.client.force_login(self.registrador)
        response = self.client.get(reverse("crear_actividad"))
        self.assertEqual(response.status_code, 302)

    def test_exportacion_de_formato_no_permitido_no_rompe_la_respuesta(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("exportar_reportes", args=["pdf"]))
        self.assertIn(response.status_code, (200, 400, 404))


@override_settings(EMAIL_BACKEND="gestion.tests_extremos.BackendCorreoFallido")
class CorreosFallidosExtremosTests(BaseSistemaTestCase):
    def test_fallo_smtp_se_registra_como_notificacion_fallida(self):
        self.actividad.carreras.set([self.carrera])
        self.actividad.jornadas.set([self.jornada])
        self.client.force_login(self.creador)
        response = self.client.post(
            reverse("enviar_invitaciones"),
            {"actividad_id": self.actividad.pk},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200)
        datos = response.json()
        self.assertFalse(datos["success"])
        self.assertEqual(datos["enviados"], 0)
        self.assertEqual(datos["fallidos"], 1)
        self.assertEqual(
            NotificacionCorreo.objects.filter(
                actividad=self.actividad, estado_envio="FALLO"
            ).count(),
            1,
        )
        self.assertEqual(len(mail.outbox), 0)


class LimitesDeInvitacionesTests(BaseSistemaTestCase):
    @override_settings(
        EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend",
        MAX_INVITACIONES_POR_ENVIO=1,
    )
    def test_se_respeta_el_limite_de_invitaciones_por_envio(self):
        self.actividad.carreras.set([self.carrera, self.carrera_2])
        self.actividad.jornadas.set([self.jornada, self.jornada_2])
        self.client.force_login(self.creador)
        response = self.client.post(
            reverse("enviar_invitaciones"),
            {"actividad_id": self.actividad.pk},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["enviados"], 1)
        self.assertEqual(
            NotificacionCorreo.objects.filter(actividad=self.actividad).count(),
            1,
        )


class IntegridadDeExportacionesTests(BaseSistemaTestCase):
    def test_excel_de_detalle_se_puede_abrir(self):
        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("exportar_reportes_detalle", args=["excel"])
        )
        libro = load_workbook(filename=BytesIO(response.content), read_only=True)
        self.assertGreaterEqual(len(libro.sheetnames), 1)

    def test_reporte_filtrado_por_actividad_responde_correctamente(self):
        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("reportes"), {"actividad": self.actividad.pk}
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["actividad_seleccionada"], str(self.actividad.pk))

    def test_auditoria_guarda_datos_de_peticion(self):
        self.client.force_login(self.admin)
        inicio = timezone.localtime(timezone.now() + timedelta(days=7))
        self.client.post(
            reverse("crear_actividad"),
            {
                "titulo": "Actividad auditada",
                "descripcion": "Descripción auditada",
                "tipo": "MASIVA",
                "lugar": "Sala auditoría",
                "fecha_inicio": inicio.strftime("%Y-%m-%dT%H:%M"),
                "fecha_fin": (inicio + timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M"),
                "carreras": [self.carrera.pk],
                "jornadas": [self.jornada.pk],
            },
            HTTP_USER_AGENT="TestBrowser/1.0",
        )
        log = LogAuditoria.objects.filter(
            usuario_sistema=self.admin, accion="Crear Actividad"
        ).first()
        self.assertIsNotNone(log)
        self.assertEqual(log.user_agent, "TestBrowser/1.0")
